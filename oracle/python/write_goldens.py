#!/usr/bin/env python3
"""Write the canonical workbook with openpyxl and XlsxWriter (golden fixtures for the Python ecosystem).

Usage: write_goldens.py <canonical.json> <out-dir>   -> <out-dir>/openpyxl/canonical.xlsx, <out-dir>/xlsxwriter/canonical.xlsx
Each golden gets a .features.json sidecar listing which structural features were applied / skipped.
"""
import datetime as dt
import json
import sys
from pathlib import Path

EXCEL_MAX = 32767
SUFFIX = "...(truncated)"


def truncate(value):
    if isinstance(value, str) and len(value) > EXCEL_MAX:
        return value[: EXCEL_MAX - len(SUFFIX)] + SUFFIX
    return value


def parse_iso(value):
    if "$date" in value:
        return dt.date.fromisoformat(value["$date"])
    if "$datetime" in value:
        return dt.datetime.fromisoformat(value["$datetime"])
    if "$time" in value:
        return dt.time.fromisoformat(value["$time"])
    return None


def is_typed(value, key):
    return isinstance(value, dict) and key in value


def sidecar(out_path: Path, generator: str, applied, skipped):
    out_path.with_suffix(".features.json").write_text(json.dumps({"generator": generator, "applied": applied, "skipped": skipped}, indent=2) + "\n")


def write_openpyxl(canonical, out_path: Path):
    import openpyxl
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    applied, skipped = [], []
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in canonical["sheets"]:
        ws = wb.create_sheet(sheet["name"])
        if sheet["hidden"]:
            ws.sheet_state = "hidden"
            applied.append("hiddenSheet")
        header = sheet["rows"][0]
        for r, row in enumerate(sheet["rows"], start=1):
            for c, value in enumerate(row, start=1):
                if value is None:
                    continue
                try:
                    if is_typed(value, "$formula"):
                        ws.cell(r, c, "=" + value["$formula"])  # openpyxl cannot write cached values
                    elif is_typed(value, "$error"):
                        ws.cell(r, c, value["$error"])  # openpyxl types '#N/A' strings as error cells
                    elif isinstance(value, dict):
                        cell = ws.cell(r, c, parse_iso(value))
                        fmt = sheet.get("numFmts", {}).get(header[c - 1]) if c - 1 < len(header) else None
                        if fmt:
                            cell.number_format = fmt
                    else:
                        cell = ws.cell(r, c, truncate(value))
                        if r == 1 and sheet.get("numFmts"):
                            cell.font = Font(bold=True)
                        elif r > 1 and isinstance(value, (int, float)) and not isinstance(value, bool):
                            fmt = sheet.get("numFmts", {}).get(header[c - 1]) if c - 1 < len(header) else None
                            if fmt:
                                cell.number_format = fmt
                except Exception as error:  # e.g. IllegalCharacterError for control characters
                    skipped.append({"feature": f"cell {sheet['name']}!r{r}c{c}", "reason": f"{type(error).__name__}: {error}"})
        if sheet.get("numFmts"):
            skipped.append({"feature": "formula cached values", "reason": "openpyxl never writes cached formula results"})
        features = sheet.get("features")
        if features:
            def attempt(name, fn):
                try:
                    fn()
                    applied.append(name)
                except Exception as error:
                    skipped.append({"feature": name, "reason": f"{type(error).__name__}: {error}"})

            attempt("merges", lambda: ws.merge_cells("A1:C1"))
            attempt("freeze", lambda: setattr(ws, "freeze_panes", "A3"))
            attempt("autoFilter", lambda: setattr(ws.auto_filter, "ref", features["autoFilter"]))
            attempt("hyperlink", lambda: setattr(ws[features["hyperlink"]["cell"]], "hyperlink", features["hyperlink"]["url"]))
            attempt("richText", lambda: ws.__setitem__(features["richText"]["cell"], CellRichText(TextBlock(InlineFont(b=True), "bold"), " and plain")))
            attempt("note", lambda: setattr(ws[features["note"]["cell"]], "comment", Comment(features["note"]["text"], "Jetstream")))

            def validation():
                dv = DataValidation(type="list", formula1='"' + ",".join(features["validation"]["list"]) + '"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(features["validation"]["range"])

            attempt("validation", validation)
            attempt(
                "conditionalFormat",
                lambda: ws.conditional_formatting.add(
                    features["conditionalFormat"]["range"],
                    CellIsRule(operator="greaterThan", formula=[str(features["conditionalFormat"]["value"])], fill=PatternFill("solid", fgColor=features["conditionalFormat"]["fillColor"][2:])),
                ),
            )
            attempt("hiddenRows", lambda: [setattr(ws.row_dimensions[n], "hidden", True) for n in features["hiddenRows"]])
            attempt("hiddenColumns", lambda: [setattr(ws.column_dimensions[letter], "hidden", True) for letter in features["hiddenColumns"]])
            attempt("columnWidths", lambda: [setattr(ws.column_dimensions[letter], "width", width) for letter, width in features["columnWidths"].items()])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    sidecar(out_path, f"openpyxl {openpyxl.__version__}", applied, skipped)


def write_xlsxwriter(canonical, out_path: Path):
    import xlsxwriter

    applied, skipped = [], []
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = xlsxwriter.Workbook(str(out_path), {"remove_timezone": True})
    bold = wb.add_format({"bold": True})
    formats = {}

    def fmt(code):
        if code not in formats:
            formats[code] = wb.add_format({"num_format": code})
        return formats[code]

    for sheet in canonical["sheets"]:
        ws = wb.add_worksheet(sheet["name"])
        if sheet["hidden"]:
            ws.hide()
            applied.append("hiddenSheet")
        header = sheet["rows"][0]
        for r, row in enumerate(sheet["rows"]):
            for c, value in enumerate(row):
                if value is None:
                    continue
                try:
                    col_name = header[c] if c < len(header) else None
                    col_fmt = fmt(sheet["numFmts"][col_name]) if r > 0 and col_name and sheet.get("numFmts", {}).get(col_name) else None
                    if is_typed(value, "$formula"):
                        cached = value.get("cached")
                        cached_text = cached["$error"] if is_typed(cached, "$error") else cached
                        ws.write_formula(r, c, "=" + value["$formula"], None, cached_text)
                    elif is_typed(value, "$error"):
                        ws.write_string(r, c, value["$error"])  # XlsxWriter has no error-cell writer; note the gap
                        skipped.append({"feature": f"error cell r{r+1}c{c+1}", "reason": "XlsxWriter has no write_error; written as text"})
                    elif isinstance(value, dict):
                        ws.write_datetime(r, c, parse_iso(value), col_fmt or fmt("yyyy-mm-dd hh:mm:ss"))
                    elif isinstance(value, bool):
                        ws.write_boolean(r, c, value)
                    elif isinstance(value, (int, float)):
                        ws.write_number(r, c, value, col_fmt)
                    elif value == "":
                        ws.write_blank(r, c, None)
                    else:
                        ws.write_string(r, c, truncate(value), bold if r == 0 and sheet.get("numFmts") else None)
                except Exception as error:
                    skipped.append({"feature": f"cell {sheet['name']}!r{r+1}c{c+1}", "reason": f"{type(error).__name__}: {error}"})
        features = sheet.get("features")
        if features:
            def attempt(name, fn):
                try:
                    fn()
                    applied.append(name)
                except Exception as error:
                    skipped.append({"feature": name, "reason": f"{type(error).__name__}: {error}"})

            attempt("merges", lambda: ws.merge_range("A1:C1", "Merged header"))
            attempt("freeze", lambda: ws.freeze_panes(2, 0))
            attempt("autoFilter", lambda: ws.autofilter(features["autoFilter"]))
            attempt("hyperlink", lambda: ws.write_url(features["hyperlink"]["cell"], features["hyperlink"]["url"], None, "Jetstream"))
            attempt("richText", lambda: ws.write_rich_string(features["richText"]["cell"], bold, "bold", " and plain"))
            attempt("note", lambda: ws.write_comment(features["note"]["cell"], features["note"]["text"], {"author": "Jetstream"}))
            attempt("validation", lambda: ws.data_validation(features["validation"]["range"], {"validate": "list", "source": features["validation"]["list"]}))
            attempt(
                "conditionalFormat",
                lambda: ws.conditional_format(
                    features["conditionalFormat"]["range"],
                    {"type": "cell", "criteria": ">", "value": features["conditionalFormat"]["value"], "format": wb.add_format({"bg_color": "#" + features["conditionalFormat"]["fillColor"][2:]})},
                ),
            )
            attempt("hiddenRows", lambda: [ws.set_row(n - 1, None, None, {"hidden": True}) for n in features["hiddenRows"]])
            attempt("hiddenColumns", lambda: [ws.set_column(f"{letter}:{letter}", None, None, {"hidden": True}) for letter in features["hiddenColumns"]])
            attempt("columnWidths", lambda: [ws.set_column(f"{letter}:{letter}", width) for letter, width in features["columnWidths"].items()])
    wb.close()
    sidecar(out_path, f"XlsxWriter {xlsxwriter.__version__}", applied, skipped)


if __name__ == "__main__":
    canonical_path, out_dir = Path(sys.argv[1]), Path(sys.argv[2])
    canonical = json.loads(canonical_path.read_text())
    write_openpyxl(canonical, out_dir / "openpyxl" / "canonical.xlsx")
    write_xlsxwriter(canonical, out_dir / "xlsxwriter" / "canonical.xlsx")
    print("wrote openpyxl + xlsxwriter goldens to", out_dir)
