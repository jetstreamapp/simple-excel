#!/usr/bin/env python3
"""Typed dump of an xlsx as the Python ecosystem sees it (same JSON shape as oracle/lib/typed.mjs).

Usage: read_dump.py <file.xlsx> --reader openpyxl|calamine [--out dump.json]
"""
import argparse
import datetime as dt
import json
import math
import sys
from importlib.metadata import version as pkg_version
from pathlib import Path


def typed_temporal(value):
    if isinstance(value, dt.datetime):
        if value.year == 1899 and value.month == 12 and value.day == 30:
            return {"$time": value.strftime("%H:%M:%S.") + f"{value.microsecond // 1000:03d}"}
        if value.hour == 0 and value.minute == 0 and value.second == 0 and value.microsecond == 0:
            return {"$date": value.strftime("%Y-%m-%d")}
        return {"$datetime": value.strftime("%Y-%m-%dT%H:%M:%S.") + f"{value.microsecond // 1000:03d}"}
    if isinstance(value, dt.date):
        return {"$date": value.strftime("%Y-%m-%d")}
    if isinstance(value, dt.time):
        return {"$time": value.strftime("%H:%M:%S.") + f"{value.microsecond // 1000:03d}"}
    if isinstance(value, dt.timedelta):
        total = value.total_seconds()
        hours, rem = divmod(int(total), 3600)
        minutes, seconds = divmod(rem, 60)
        ms = int(round((total - int(total)) * 1000))
        return {"$time": f"{hours:02d}:{minutes:02d}:{seconds:02d}.{ms:03d}"}
    return None


ERROR_CODES = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#GETTING_DATA"}


def number(value):
    """JSON cannot carry inf/nan; keep them visible as typed markers."""
    if isinstance(value, float) and not math.isfinite(value):
        return {"$number": str(value)}
    return value


def normalize_rows(rows):
    trimmed = []
    for row in rows:
        end = len(row)
        while end > 0 and row[end - 1] is None:
            end -= 1
        trimmed.append(row[:end])
    while trimmed and not trimmed[-1]:
        trimmed.pop()
    return trimmed


def dump_openpyxl(path):
    import openpyxl
    from openpyxl.cell.rich_text import CellRichText

    wb = openpyxl.load_workbook(path, data_only=False, rich_text=True)
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows():
            out = []
            for cell in row:
                value = cell.value
                if value is None:
                    out.append(None)
                elif isinstance(value, CellRichText):
                    out.append(str(value))
                elif isinstance(value, str) and value.startswith("=") and cell.data_type == "f":
                    out.append({"$formula": value[1:], "cached": None})
                elif cell.data_type == "e" or (isinstance(value, str) and value in ERROR_CODES):
                    out.append({"$error": value})
                elif temporal := typed_temporal(value):
                    out.append(temporal)
                elif isinstance(value, bool):
                    out.append(value)
                elif isinstance(value, (int, float)):
                    out.append(number(value))
                else:
                    out.append(str(value))
            rows.append(out)
        sheets.append(
            {
                "name": ws.title,
                "hidden": ws.sheet_state != "visible",
                "rows": normalize_rows(rows),
                "merges": [str(r) for r in ws.merged_cells.ranges],
                "autoFilter": ws.auto_filter.ref or None,
                "freeze": ws.freeze_panes,
                "hiddenRows": [n for n, d in ws.row_dimensions.items() if d.hidden],
                "hiddenColumns": [k for k, d in ws.column_dimensions.items() if d.hidden],
                "hyperlinks": [{"cell": h.ref, "url": h.target} for h in ws._hyperlinks] if hasattr(ws, "_hyperlinks") else [],
                "comments": [{"cell": c.coordinate, "text": c.comment.text} for row in ws.iter_rows() for c in row if c.comment],
            }
        )
    return f"openpyxl {pkg_version('openpyxl')}", sheets


def dump_calamine(path):
    import python_calamine
    from python_calamine import CalamineWorkbook

    wb = CalamineWorkbook.from_path(path)
    sheets = []
    for meta in wb.sheets_metadata:
        ws = wb.get_sheet_by_name(meta.name)
        rows = []
        for row in ws.to_python(skip_empty_area=False):
            out = []
            for value in row:
                if value == "" or value is None:
                    out.append(None if value is None else "")
                elif temporal := typed_temporal(value):
                    out.append(temporal)
                elif isinstance(value, bool):
                    out.append(value)
                elif isinstance(value, (int, float)):
                    out.append(number(value))
                elif isinstance(value, str) and value in ERROR_CODES:
                    out.append({"$error": value})
                else:
                    out.append(str(value))
            rows.append(out)
        sheets.append({"name": meta.name, "hidden": str(meta.visible) != "SheetVisibleEnum.Visible" and meta.visible != "visible", "rows": normalize_rows(rows)})
    return f"python-calamine {pkg_version('python-calamine')}", sheets


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("file")
    parser.add_argument("--reader", choices=["openpyxl", "calamine"], required=True)
    parser.add_argument("--out")
    args = parser.parse_args()
    try:
        reader, sheets = dump_openpyxl(args.file) if args.reader == "openpyxl" else dump_calamine(args.file)
        dump = {"reader": reader, "file": args.file, "sheets": sheets}
    except Exception as error:  # a reader that cannot open the file is a result, not a crash
        dump = {"reader": args.reader, "file": args.file, "sheets": [], "error": f"{type(error).__name__}: {error}"}
    text = json.dumps(dump, indent=2, ensure_ascii=False) + "\n"
    if args.out:
        Path(args.out).write_text(text)
        print(f"wrote {args.out} ({len(dump['sheets'])} sheets)" + (f" ERROR {dump['error']}" if "error" in dump else ""))
    else:
        sys.stdout.write(text)


if __name__ == "__main__":
    main()
